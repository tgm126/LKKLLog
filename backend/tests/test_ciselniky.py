import pytest
from openpyxl import load_workbook

from lety.ciselniky import ChybaNacteni, nacti, vytvor_sablonu
from lety.models import Letadlo, Letiste, Ucel, Uloha
from osoby.models import Kategorie, Opravneni, Osoba, Uroven

pytestmark = pytest.mark.django_db


@pytest.fixture
def vyplnena_sablona(tmp_path):
    cesta = tmp_path / "ciselniky.xlsx"
    vytvor_sablonu(cesta)
    wb = load_workbook(cesta)
    wb["Osoby"].append(
        ["Jan", "Novák", "Novak@Example.com", "ano", None, None, None, None, None, "731 123 456"]
    )
    wb["Osoby"].append(["Karel", "Cizí", None, None, None, None, "ano", None, None])
    wb["Osoby"].append(
        ["Test", "Pilot", "tomas+pilot@example.com", None, None, None, None, None, "ano"]
    )
    wb["Oprávnění"].append(["Jan", "Novák", "Kluzák", "Instruktor", "31.12.2027"])
    wb["Oprávnění"].append(["Karel", "Cizí", "Motorové", "Examinátor", None])
    wb["Letadla"].append(["ok-0815", "L-13 Blaník", "Kluzák", 2, None, None, None, None, 1])
    wb["Letadla"].append(["OK-ABC", "Z-226", "Motorové", 2, 150, "ano", None, None, 2])
    wb["Úlohy"].append(["ZVP kluzáky", "12", "Okruhy", "Výcvik, Výcvik sólo", "Kluzák", 12])
    wb["Úlohy"].append(
        [
            "Přezkoušení",
            "PZ-1",
            "Přezkoušení odborné způsobilosti",
            "Přezkoušení",
            "Kluzák, Motorové, TMG",
            None,
        ]
    )
    wb.save(cesta)
    return cesta


def test_nacteni_ciselniku(vyplnena_sablona):
    vysledek = nacti(vyplnena_sablona)

    assert vysledek.zalozeno == {"Osoby": 3, "Oprávnění": 2, "Letadla": 2, "Letiště": 2, "Úlohy": 2}
    novak = Osoba.objects.get(email="novak@example.com")
    assert novak.role_casomeric and not novak.has_usable_password()
    assert Osoba.objects.get(prijmeni="Cizí").externi
    assert Osoba.objects.get(prijmeni="Pilot").testovaci
    assert not novak.testovaci
    assert novak.telefon == "+420731123456"
    assert Opravneni.objects.get(osoba=novak).uroven == Uroven.INSTRUKTOR
    assert Letadlo.objects.get(imatrikulace="OK-0815").kategorie == Kategorie.KLUZAK
    assert Letadlo.objects.get(imatrikulace="OK-ABC").vlecne
    assert Letiste.objects.get(domovske=True).icao == "LKKL"
    uloha = Uloha.objects.get(kod="12")
    assert uloha.ucely == [Ucel.VYCVIK, Ucel.VYCVIK_SOLO]


def test_opakovane_nacteni_jen_aktualizuje(vyplnena_sablona):
    nacti(vyplnena_sablona)
    vysledek = nacti(vyplnena_sablona)
    assert vysledek.zalozeno == {}
    assert Osoba.objects.count() == 3


def test_pri_chybe_se_nenacte_nic(vyplnena_sablona):
    wb = load_workbook(vyplnena_sablona)
    wb["Letadla"].append(["OK-XYZ", "Neznámý", "Vzducholoď", 2])
    wb["Oprávnění"].append(["Neexistující", "Člověk", "Kluzák", "Pilot"])
    wb.save(vyplnena_sablona)

    with pytest.raises(ChybaNacteni) as chyba:
        nacti(vyplnena_sablona)

    assert len(chyba.value.chyby) == 2
    assert "Letadla, řádek 4" in chyba.value.chyby[1]
    assert Osoba.objects.count() == 0
    assert Letadlo.objects.count() == 0
