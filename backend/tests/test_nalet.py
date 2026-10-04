import io
from datetime import timedelta

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from lety.models import Letadlo, StavLetu
from osoby.models import Kategorie

from .pomocne import let_v

pytestmark = pytest.mark.django_db


@pytest.fixture
def lety(svet):
    """Lety pilota za poslední dny: PIC, žák, soukromé letadlo, zrušený, dozor, cizí."""
    ted = timezone.now().replace(microsecond=0)
    den = ted - timedelta(days=2)
    let_v(svet, den, minut=60, pocet_tg=2)  # PIC, motor
    vycvik = let_v(svet, den + timedelta(hours=2), minut=30, pilot=svet.instruktor)
    vycvik.posadka.create(osoba=svet.pilot, funkce="zak")
    soukrome = Letadlo.objects.create(
        imatrikulace="OK-SOU", typ="Soukromý", kategorie=Kategorie.UL, soukrome=True
    )
    let_v(svet, den + timedelta(hours=4), minut=20, letadlo=soukrome, soukrome=True)
    let_v(svet, den + timedelta(hours=6), stav=StavLetu.ZRUSEN, duvod_zruseni="omyl")
    dozor = let_v(svet, den + timedelta(hours=8), minut=40, pilot=svet.zak, letadlo=svet.kluzak)
    dozor.posadka.create(osoba=svet.pilot, funkce="dozor")
    let_v(svet, den + timedelta(hours=10), pilot=svet.cizi_pilot)
    # Let mimo období (před rokem a kousek).
    let_v(svet, ted - timedelta(days=400), minut=90)
    return svet


def test_souhrn_naletu(jako, lety):
    od = (timezone.now() - timedelta(days=30)).date()
    data = jako(lety.pilot).get(f"/api/nalet?od={od}").json()
    assert data["souhrn"]["celkem"] == {"lety": 3, "minuty": 60 + 30 + 20, "pristani": 3 + 1 + 1}
    radky = {(r["kategorie"], r["funkce"]): r["minuty"] for r in data["souhrn"]["podle_kategorie"]}
    assert radky == {("Motorové", "PIC"): 60, ("Motorové", "Žák"): 30, ("UL", "PIC"): 20}
    assert [let["moje_funkce"] for let in data["lety"]] == ["pic", "zak", "pic"][::-1]
    starty = {(s["kategorie"], s["zpusob"]): s["pocet"] for s in data["souhrn"]["starty"]}
    assert starty == {("Motorové", "Vlastní"): 2, ("UL", "Vlastní"): 1}


def test_nalet_jen_vlastni_a_za_obdobi(jako, lety):
    od = (timezone.now() - timedelta(days=30)).date()
    cizi = jako(lety.cizi_pilot).get(f"/api/nalet?od={od}").json()
    assert cizi["souhrn"]["celkem"]["lety"] == 1
    # Delší období zahrne i let před rokem.
    od = (timezone.now() - timedelta(days=500)).date()
    data = jako(lety.pilot).get(f"/api/nalet?od={od}").json()
    assert data["souhrn"]["celkem"]["lety"] == 4
    assert jako(lety.pilot).get("/api/nalet?od=2026-05-01&do=2026-04-01").status_code == 400


def test_export_naletu(jako, lety):
    od = (timezone.now() - timedelta(days=30)).date()
    odpoved = jako(lety.pilot).get(f"/api/nalet/export.xlsx?od={od}")
    assert odpoved.status_code == 200
    kniha = load_workbook(io.BytesIO(odpoved.content))
    assert kniha.sheetnames == ["Lety", "Souhrn", "Parametry"]
    radky = list(kniha["Lety"].values)
    assert len(radky) == 1 + 3
    vycvik = dict(zip(radky[0], radky[2], strict=True))
    assert vycvik["Funkce"] == "Žák" and vycvik["Další posádka"] == "Test Instruktor (PIC)"
    assert list(kniha["Souhrn"].values)[-1] == ("Celkem", None, 3, 110, '1°50"', 5)
