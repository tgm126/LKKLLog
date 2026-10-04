import io
from datetime import UTC, datetime, time, timedelta

import pytest
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from lety import uzaverky
from lety.models import AuditLog, Let, StavLetu, Uzaverka
from osoby.models import Osoba

from .pomocne import let_v, normalni, post
from .test_opravy import oprava

pytestmark = pytest.mark.django_db


@pytest.fixture
def jako():
    """Každá osoba má vlastního klienta (sdílený klient by přihlášení přepsal)."""

    def prihlasit(osoba):
        klient = Client()
        klient.force_login(osoba)
        return klient

    return prihlasit


def v(den, hodina, minuta=0):
    return datetime.combine(den, time(hodina, minuta), tzinfo=UTC)


@pytest.fixture
def den():
    return timezone.now().date() - timedelta(days=3)


@pytest.fixture
def minuly_mesic():
    return (timezone.now().date().replace(day=1) - timedelta(days=1)).replace(day=1)


def uzavrit(klient, typ, obdobi):
    return post(klient, "/api/uzaverky", {"typ": typ, "obdobi": obdobi.isoformat()})


def detail_letu(klient, let):
    return klient.get(f"/api/lety/{let.pk}").json()


def test_den_nejde_uzavrit_dokud_neco_leti(jako, svet, den):
    let_v(svet, v(den, 9), minut=60, pocet_tg=2)
    ve_vzduchu = let_v(svet, v(den, 11), letadlo=svet.dvoumistne, stav=StavLetu.VE_VZDUCHU)
    klient = jako(svet.casomeric)

    odpoved = uzavrit(klient, "den", den)
    assert odpoved.status_code == 409
    assert odpoved.json()["kod"] == "neukonceno"
    assert "OK-TVA (ve vzduchu)" in odpoved.json()["detail"]

    ve_vzduchu.stav = StavLetu.UKONCEN
    ve_vzduchu.cas_pristani = v(den, 11, 30)
    ve_vzduchu.misto_pristani = svet.lkkl
    ve_vzduchu.save()
    odpoved = uzavrit(klient, "den", den)
    assert odpoved.status_code == 200, odpoved.json()
    assert odpoved.json()["verze"] == 1

    u = Uzaverka.objects.get()
    assert u.souhrn["celkem"] == {"lety": 2, "minuty": 90, "tg": 2, "navijak": 0, "vlek": 0}
    assert u.souhrn["starty"]["vlastni"] == 2
    assert u.souhrn["podle_osob"][0]["minuty"] == 90
    assert AuditLog.objects.filter(akce="uzaverka", objekt_id=u.pk).exists()


def test_kdo_smi_uzavrit_den(jako, svet, den):
    let_v(svet, v(den, 9))
    assert uzavrit(jako(svet.pilot), "den", den).status_code == 403
    zitra = timezone.now().date() + timedelta(days=1)
    assert uzavrit(jako(svet.casomeric), "den", zitra).status_code == 400
    assert uzavrit(jako(svet.ucetni), "den", den).status_code == 200


def test_pripraveny_let_brani_uzavreni_dneska(jako, svet):
    dnes = timezone.now().date()
    post(jako(svet.pilot), "/api/lety", normalni(svet, akce="pripravit"))
    odpoved = uzavrit(jako(svet.casomeric), "den", dnes)
    assert odpoved.status_code == 409
    assert "připraven" in odpoved.json()["detail"]


def test_v_uzavrenem_dni_pilot_jen_cte(jako, svet, den):
    let = let_v(svet, v(den, 9))
    uzavrit(jako(svet.casomeric), "den", den)

    pilot = jako(svet.pilot)
    data = detail_letu(pilot, let)
    assert data["muze_ovladat"] is False
    odpoved = post(pilot, f"/api/lety/{let.pk}/oprava", oprava(data, pocet_tg=1))
    assert odpoved.status_code == 403
    assert odpoved.json()["kod"] == "uzavreno"
    assert post(pilot, f"/api/lety/{let.pk}/zrusit", {"duvod": "omyl"}).status_code == 403

    # Dopsat let do uzavřeného dne smí jen časoměřič nebo účetní.
    dopsany = normalni(
        svet,
        letadlo_id=svet.dvoumistne.pk,
        akce="dopsat",
        cas_vzletu=v(den, 14).isoformat(),
        cas_pristani=v(den, 14, 30).isoformat(),
    )
    assert post(pilot, "/api/lety", dopsany).status_code == 403
    odpoved = post(jako(svet.casomeric), "/api/lety", dopsany)
    assert odpoved.status_code == 200
    assert odpoved.json()["opraveno_po_uzaverce"] is True


def test_oprava_po_uzaverce_se_oznaci_a_lze_prepocitat(jako, svet, den):
    let = let_v(svet, v(den, 9), minut=30)
    casomeric = jako(svet.casomeric)
    uzavrit(casomeric, "den", den)

    data = detail_letu(casomeric, let)
    zmena = oprava(data, cas_pristani=v(den, 9, 45).isoformat())
    odpoved = post(casomeric, f"/api/lety/{let.pk}/oprava", zmena)
    assert odpoved.status_code == 200, odpoved.json()
    assert odpoved.json()["opraveno_po_uzaverce"] is True

    prehled = casomeric.get(f"/api/prehled?den={den}").json()
    assert prehled["uzaverka"]["uzaverka"]["verze"] == 1
    assert prehled["uzaverka"]["zmeny"] == 1

    detail = casomeric.get(f"/api/uzaverky/detail?typ=den&obdobi={den}").json()
    assert detail["rozdil"]["celkem"]["minuty"] == 15
    assert [z["akce"] for z in detail["lety"][0]["zaznamy"]] == ["Oprava"]
    assert detail["lety"][0]["zaznamy"][0]["duvod"]  # čitelný důvod

    # Přepočet = nová verze; stará zůstává v historii.
    odpoved = uzavrit(casomeric, "den", den)
    assert odpoved.json()["verze"] == 2
    detail = casomeric.get(f"/api/uzaverky/detail?typ=den&obdobi={den}").json()
    assert [u["verze"] for u in detail["verze"]] == [2, 1]
    assert detail["rozdil"] is None
    assert detail["souhrn"]["celkem"]["minuty"] == 45


def test_presun_letu_do_uzavreneho_dne_hlida_prava(jako, svet, den):
    let = let_v(svet, v(den + timedelta(days=1), 9))
    uzavrit(jako(svet.casomeric), "den", den)
    pilot = jako(svet.pilot)
    data = detail_letu(pilot, let)
    assert data["muze_ovladat"] is True
    zmena = oprava(data, cas_vzletu=v(den, 9).isoformat(), cas_pristani=v(den, 9, 30).isoformat())
    assert post(pilot, f"/api/lety/{let.pk}/oprava", zmena).status_code == 403


def test_mesicni_uzaverka(jako, svet, minuly_mesic):
    prvni = let_v(svet, v(minuly_mesic.replace(day=5), 9), minut=40)
    let_v(svet, v(minuly_mesic.replace(day=6), 10), minut=20, plati_aeroklub=True, platce=None)
    ucetni = jako(svet.ucetni)
    casomeric = jako(svet.casomeric)

    assert uzavrit(casomeric, "mesic", minuly_mesic).status_code == 403
    odpoved = uzavrit(ucetni, "mesic", minuly_mesic)
    assert odpoved.status_code == 409
    assert odpoved.json()["kod"] == "neuzavrene_dny"
    assert "5." in odpoved.json()["detail"]
    assert uzavrit(ucetni, "mesic", timezone.now().date()).status_code == 403  # ještě běží

    for d in (5, 6):
        assert uzavrit(casomeric, "den", minuly_mesic.replace(day=d)).status_code == 200
    odpoved = uzavrit(ucetni, "mesic", minuly_mesic.replace(day=17))
    assert odpoved.status_code == 200
    assert odpoved.json()["obdobi"] == minuly_mesic.isoformat()

    # V uzavřeném měsíci časoměřič jen čte, účetní opravuje ukončené lety.
    data = detail_letu(casomeric, prvni)
    assert data["muze_ovladat"] is False
    zmena = oprava(data, pocet_tg=2)
    odpoved = post(casomeric, f"/api/lety/{prvni.pk}/oprava", zmena)
    assert odpoved.status_code == 403
    assert "účetní" in odpoved.json()["detail"]
    assert uzavrit(casomeric, "den", minuly_mesic.replace(day=5)).status_code == 403

    data = detail_letu(ucetni, prvni)
    assert data["muze_ovladat"] is True
    zmena = oprava(data, cas_pristani=(prvni.cas_pristani + timedelta(minutes=5)).isoformat())
    odpoved = post(ucetni, f"/api/lety/{prvni.pk}/oprava", zmena)
    assert odpoved.status_code == 200, odpoved.json()

    mesic = ucetni.get(f"/api/uzaverky?mesic={minuly_mesic}").json()
    assert mesic["uzaverka"]["verze"] == 1
    assert mesic["zmeny"] == 1
    assert [d["zmeny"] for d in mesic["dny"]] == [0, 1]  # nejnovější den první
    assert mesic["skoncil"] is True


def test_admin_znovu_otevre_uzaverku(jako, svet, den):
    let = let_v(svet, v(den, 9))
    uzavrit(jako(svet.casomeric), "den", den)
    admin = Osoba.objects.create_superuser("admin@example.com", "x")
    assert uzaverky.znovu_otevrit("den", den, admin) == 1

    pilot = jako(svet.pilot)
    assert detail_letu(pilot, let)["muze_ovladat"] is True
    prehled = pilot.get(f"/api/prehled?den={den}").json()
    assert prehled["uzaverka"]["uzaverka"] is None
    assert uzavrit(jako(svet.casomeric), "den", den).json()["verze"] == 2


def test_nahled_a_export(jako, svet, den):
    let_v(svet, v(den, 9), minut=75)
    casomeric = jako(svet.casomeric)
    nahled = casomeric.get(f"/api/uzaverky/nahled?typ=den&obdobi={den}").json()
    assert nahled["lze"] is True
    assert nahled["souhrn"]["celkem"]["minuty"] == 75
    assert nahled["posledni"] is None
    assert jako(svet.pilot).get(f"/api/uzaverky/nahled?typ=den&obdobi={den}").json()["lze"] is False

    u = uzavrit(casomeric, "den", den).json()
    assert casomeric.get(f"/api/uzaverky/{u['id']}/export.xlsx").status_code == 403
    odpoved = jako(svet.ucetni).get(f"/api/uzaverky/{u['id']}/export.xlsx")
    assert odpoved.status_code == 200
    kniha = load_workbook(io.BytesIO(odpoved.content))
    assert kniha.sheetnames == ["Uzávěrka", "Podle letadel", "Podle plátců", "Podle osob", "Starty"]
    radky = list(kniha["Podle letadel"].values)
    assert radky[-1][:6] == ("Celkem", None, None, 1, 75, '1°15"')


def test_lety_po_uzaverce_v_exportu_vypisu(jako, svet, den):
    let = let_v(svet, v(den, 9))
    Let.objects.filter(pk=let.pk).update(opraveno_po_uzaverce=True)
    odpoved = jako(svet.ucetni).get(f"/api/vypis/export.xlsx?od={den}&do={den}")
    radky = list(load_workbook(io.BytesIO(odpoved.content))["Lety"].values)
    sloupec = radky[0].index("Opraveno po uzávěrce")
    assert radky[1][sloupec] == "ano"


def test_akce_znovu_otevrit_v_administraci(client, svet, den):
    let_v(svet, v(den, 9))
    uzaverky.uzavrit("den", den, svet.casomeric)
    admin = Osoba.objects.create_superuser("admin@example.com", "x")
    client.force_login(admin)
    u = Uzaverka.objects.get()
    odpoved = client.post(
        "/admin/lety/uzaverka/",
        {"action": "znovu_otevrit", "_selected_action": [u.pk]},
        follow=True,
    )
    assert "Znovu otevřeno období: 1." in odpoved.text
    u.refresh_from_db()
    assert u.znovu_otevreno is not None
    assert AuditLog.objects.filter(akce="otevreni").count() == 1
