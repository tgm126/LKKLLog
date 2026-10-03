"""Úvodní načtení číselníků z Excelu (osoby, oprávnění, letadla, letiště, úlohy).

Šablonu vytvoří `manage.py sablona_ciselniku`, data načte `manage.py nacti_ciselniky`.
Načítání je opakovatelné: existující záznamy se podle klíče aktualizují, nové se založí,
nic se nemaže.
"""

from dataclasses import dataclass, field
from datetime import date, datetime

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from osoby.models import Kategorie, Opravneni, Osoba, Uroven

from .models import Letadlo, Letiste, Ucel, Uloha

ANO_NE = ["ano", "ne"]

# List → seznam sloupců (záhlaví, povinný, výběr hodnot nebo None)
LISTY: dict[str, list[tuple[str, bool, list[str] | None]]] = {
    "Osoby": [
        ("Jméno", True, None),
        ("Příjmení", True, None),
        ("E-mail", False, None),
        ("Časoměřič / věž", False, ANO_NE),
        ("Účetní", False, ANO_NE),
        ("Admin", False, ANO_NE),
        ("Externí", False, ANO_NE),
        ("Aktivní", False, ANO_NE),
    ],
    "Oprávnění": [
        ("Jméno", True, None),
        ("Příjmení", True, None),
        ("Kategorie", True, Kategorie.labels),
        ("Úroveň", True, Uroven.labels),
        ("Platné do", False, None),
    ],
    "Letadla": [
        ("Imatrikulace", True, None),
        ("Typ", True, None),
        ("Kategorie", True, Kategorie.labels),
        ("Počet míst", True, None),
        ("Max. doba letu [min]", False, None),
        ("Vlečné", False, ANO_NE),
        ("Soukromé", False, ANO_NE),
        ("Aktivní", False, ANO_NE),
        ("Pořadí", False, None),
    ],
    "Letiště": [
        ("ICAO", False, None),
        ("Název", True, None),
        ("Domovské", False, ANO_NE),
        ("Mimo letiště (terén)", False, ANO_NE),
        ("Pořadí", False, None),
    ],
    "Úlohy": [
        ("Osnova", False, None),
        ("Kód", True, None),
        ("Název", True, None),
        ("Účely", True, None),
        ("Kategorie", True, None),
        ("Pořadí", False, None),
    ],
}

NAVOD = [
    "Šablona pro úvodní načtení číselníků do LKKL Log",
    "",
    "• Každý list = jeden číselník. Řádek 1 je záhlaví, data pište od řádku 2.",
    "• Tučné sloupce jsou povinné. Kde je šipka, vybírejte z nabídky.",
    "• ano/ne: prázdné = výchozí hodnota (Aktivní = ano, ostatní = ne).",
    "• Oprávnění: osoba se dohledá podle jména a příjmení z listu Osoby.",
    "  Každé oprávnění na samostatném řádku (např. kluzák–instruktor, motorové–pilot).",
    "• Úlohy: Účely a Kategorie mohou mít více hodnot oddělených čárkou,",
    f"  účely: {', '.join(Ucel.labels[:4])}",
    f"  kategorie: {', '.join(Kategorie.labels)}",
    "  příklad: Účely = „Výcvik, Výcvik sólo“, Kategorie = „Kluzák“",
    "• Přezkoušení zapište jako úlohy s účelem „Přezkoušení“ (jedna úloha = jeden typ).",
    "• Letiště LKKL a „Mimo letiště (terén)“ jsou předvyplněná.",
    "• Načtení lze opakovat: existující záznamy se aktualizují, nic se nemaže.",
]


def vytvor_sablonu(cesta) -> None:
    wb = Workbook()
    navod = wb.active
    navod.title = "Návod"
    for i, radek in enumerate(NAVOD, start=1):
        navod.cell(row=i, column=1, value=radek)
    navod["A1"].font = Font(bold=True, size=14)
    navod.column_dimensions["A"].width = 90

    zahlavi_fill = PatternFill("solid", fgColor="DDE7F3")
    for nazev, sloupce in LISTY.items():
        ws = wb.create_sheet(nazev)
        for col, (zahlavi, povinny, vyber) in enumerate(sloupce, start=1):
            bunka = ws.cell(row=1, column=col, value=zahlavi)
            bunka.font = Font(bold=povinny)
            bunka.fill = zahlavi_fill
            bunka.alignment = Alignment(wrap_text=True, vertical="top")
            pismeno = get_column_letter(col)
            ws.column_dimensions[pismeno].width = max(14, len(zahlavi) + 4)
            if vyber:
                dv = DataValidation(
                    type="list", formula1='"' + ",".join(vyber) + '"', allow_blank=True
                )
                dv.error = "Vyberte hodnotu z nabídky."
                ws.add_data_validation(dv)
                dv.add(f"{pismeno}2:{pismeno}500")
        ws.freeze_panes = "A2"

    letiste = wb["Letiště"]
    letiste.append(["LKKL", "Kladno", "ano", "ne", 1])
    letiste.append([None, "Mimo letiště (terén)", "ne", "ano", 999])
    wb.save(cesta)


class ChybaNacteni(Exception):
    def __init__(self, chyby: list[str]):
        super().__init__("\n".join(chyby))
        self.chyby = chyby


@dataclass
class Vysledek:
    zalozeno: dict[str, int] = field(default_factory=dict)
    aktualizovano: dict[str, int] = field(default_factory=dict)

    def pricti(self, list_: str, novy: bool) -> None:
        cil = self.zalozeno if novy else self.aktualizovano
        cil[list_] = cil.get(list_, 0) + 1


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _ano_ne(v, vychozi: bool) -> bool:
    t = _text(v).lower()
    if not t:
        return vychozi
    if t in ("ano", "a", "1", "true", "x"):
        return True
    if t in ("ne", "n", "0", "false"):
        return False
    raise ValueError(f"očekávám ano/ne, ne „{v}“")


def _volba(v, choices) -> str:
    t = _text(v).lower()
    for value, label in choices:
        if t in (value.lower(), label.lower()):
            return value
    moznosti = ", ".join(popis for _, popis in choices)
    raise ValueError(f"neznámá hodnota „{v}“ (možnosti: {moznosti})")


def _volby(v, choices) -> list[str]:
    hodnoty = [_volba(x, choices) for x in _text(v).split(",") if x.strip()]
    if not hodnoty:
        raise ValueError("chybí hodnota")
    return list(dict.fromkeys(hodnoty))


def _cislo(v, vychozi: int | None = None) -> int | None:
    if _text(v) == "":
        return vychozi
    return int(float(v))


def _datum(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _text(v)
    if not t:
        return None
    return datetime.strptime(t, "%d.%m.%Y").date()


def _radky(wb, nazev):
    if nazev not in wb.sheetnames:
        return
    ws = wb[nazev]
    for cislo, radek in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if any(_text(v) for v in radek):
            yield cislo, radek


def nacti(cesta) -> Vysledek:
    """Načte číselníky. Při jakékoli chybě se nezmění nic (vše v jedné transakci)."""
    wb = load_workbook(cesta, data_only=True)
    vysledek = Vysledek()
    chyby: list[str] = []

    with transaction.atomic():
        for list_, zpracuj in [
            ("Osoby", _osoba),
            ("Oprávnění", _opravneni),
            ("Letadla", _letadlo),
            ("Letiště", _letiste),
            ("Úlohy", _uloha),
        ]:
            for cislo, radek in _radky(wb, list_):
                try:
                    with transaction.atomic():
                        vysledek.pricti(list_, zpracuj(radek))
                except Exception as e:  # noqa: BLE001 – chybu vrátíme uživateli s číslem řádku
                    chyby.append(f"{list_}, řádek {cislo}: {e}")
        if chyby:
            transaction.set_rollback(True)
            raise ChybaNacteni(chyby)
    return vysledek


def _osoba(r) -> bool:
    jmeno, prijmeni, email = _text(r[0]), _text(r[1]), _text(r[2]).lower() or None
    if not jmeno or not prijmeni:
        raise ValueError("chybí jméno nebo příjmení")
    hodnoty = {
        "jmeno": jmeno,
        "prijmeni": prijmeni,
        "email": email,
        "role_casomeric": _ano_ne(r[3], False),
        "role_ucetni": _ano_ne(r[4], False),
        "is_staff": _ano_ne(r[5], False),
        "externi": _ano_ne(r[6], False),
        "is_active": _ano_ne(r[7], True),
    }
    osoba = None
    if email:
        osoba = Osoba.objects.filter(email=email).first()
    if osoba is None:
        osoba = Osoba.objects.filter(jmeno=jmeno, prijmeni=prijmeni).first()
    novy = osoba is None
    if novy:
        osoba = Osoba(**hodnoty)
        osoba.set_unusable_password()
    else:
        for k, v in hodnoty.items():
            setattr(osoba, k, v)
    osoba.full_clean(exclude=["password"])
    osoba.save()
    return novy


def _najdi_osobu(jmeno, prijmeni) -> Osoba:
    osoby = list(Osoba.objects.filter(jmeno=_text(jmeno), prijmeni=_text(prijmeni)))
    if not osoby:
        raise ValueError(f"osoba {jmeno} {prijmeni} není v listu Osoby")
    if len(osoby) > 1:
        raise ValueError(f"osoba {jmeno} {prijmeni} je v listu Osoby vícekrát")
    return osoby[0]


def _opravneni(r) -> bool:
    osoba = _najdi_osobu(r[0], r[1])
    _, novy = Opravneni.objects.update_or_create(
        osoba=osoba,
        kategorie=_volba(r[2], Kategorie.choices),
        uroven=_volba(r[3], Uroven.choices),
        defaults={"platne_do": _datum(r[4])},
    )
    return novy


def _letadlo(r) -> bool:
    imatrikulace = _text(r[0]).upper()
    if not imatrikulace:
        raise ValueError("chybí imatrikulace")
    letadlo = Letadlo.objects.filter(imatrikulace=imatrikulace).first() or Letadlo(
        imatrikulace=imatrikulace
    )
    novy = letadlo.pk is None
    letadlo.typ = _text(r[1])
    letadlo.kategorie = _volba(r[2], Kategorie.choices)
    letadlo.pocet_mist = _cislo(r[3])
    letadlo.max_doba_min = _cislo(r[4])
    letadlo.vlecne = _ano_ne(r[5], False)
    letadlo.soukrome = _ano_ne(r[6], False)
    letadlo.aktivni = _ano_ne(r[7], True)
    letadlo.poradi = _cislo(r[8], 100)
    letadlo.full_clean()
    letadlo.save()
    return novy


def _letiste(r) -> bool:
    icao, nazev = _text(r[0]).upper() or None, _text(r[1])
    if not nazev:
        raise ValueError("chybí název")
    letiste = (
        Letiste.objects.filter(icao=icao).first()
        if icao
        else Letiste.objects.filter(icao__isnull=True, nazev=nazev).first()
    ) or Letiste(icao=icao)
    novy = letiste.pk is None
    letiste.nazev = nazev
    letiste.domovske = _ano_ne(r[2], False)
    letiste.teren = _ano_ne(r[3], False)
    letiste.poradi = _cislo(r[4], 100)
    letiste.full_clean()
    letiste.save()
    return novy


def _uloha(r) -> bool:
    osnova, kod = _text(r[0]), _text(r[1])
    if not kod:
        raise ValueError("chybí kód")
    uloha = Uloha.objects.filter(osnova=osnova, kod=kod).first() or Uloha(osnova=osnova, kod=kod)
    novy = uloha.pk is None
    uloha.nazev = _text(r[2])
    uloha.ucely = _volby(r[3], Ucel.choices)
    uloha.kategorie = _volby(r[4], Kategorie.choices)
    uloha.poradi = _cislo(r[5], 100)
    uloha.full_clean()
    uloha.save()
    return novy
