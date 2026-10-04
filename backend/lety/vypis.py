"""Výpis letů za období, souhrny a export do Excelu / CSV (etapa 7 návrhu)."""

import csv
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta

from django.db.models import QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from osoby.models import Kategorie

from .models import (
    DuvodZruseni,
    FunkcePosadky,
    KratkyLet,
    Let,
    StavLetu,
    Ucel,
    ZpusobVzletu,
)

MAX_DNU = 400  # ochrana proti obřím výpisům


class ChybaVypisu(Exception):
    pass


@dataclass
class Filtr:
    od: date
    do: date
    letadlo_id: int | None = None
    osoba_id: int | None = None  # kdokoli v posádce
    platce_id: int | None = None  # 0 = aeroklub
    kategorie: str | None = None
    ucel: str | None = None
    zpusob_vzletu: str | None = None
    vcetne_soukromych: bool = False
    vcetne_zrusenych: bool = False

    def over(self) -> None:
        if self.do < self.od:
            raise ChybaVypisu("Konec období je před začátkem.")
        if (self.do - self.od).days > MAX_DNU:
            raise ChybaVypisu(f"Období může mít nejvýš {MAX_DNU} dní.")


def lety(filtr: Filtr, zaklad: QuerySet) -> list[Let]:
    """Lety se vzletem v období (UTC), seřazené podle vzletu."""
    filtr.over()
    zacatek = datetime.combine(filtr.od, time.min, tzinfo=UTC)
    konec = datetime.combine(filtr.do + timedelta(days=1), time.min, tzinfo=UTC)
    q = zaklad.filter(cas_vzletu__gte=zacatek, cas_vzletu__lt=konec)
    if not filtr.vcetne_zrusenych:
        q = q.exclude(stav=StavLetu.ZRUSEN)
    if not filtr.vcetne_soukromych:
        q = q.filter(soukrome=False)
    if filtr.letadlo_id:
        q = q.filter(letadlo_id=filtr.letadlo_id)
    if filtr.osoba_id:
        q = q.filter(posadka__osoba_id=filtr.osoba_id)
    if filtr.platce_id == 0:
        q = q.filter(plati_aeroklub=True)
    elif filtr.platce_id:
        q = q.filter(platce_id=filtr.platce_id)
    if filtr.kategorie:
        q = q.filter(letadlo__kategorie=filtr.kategorie)
    if filtr.ucel:
        q = q.filter(ucel=filtr.ucel)
    if filtr.zpusob_vzletu:
        q = q.filter(zpusob_vzletu=filtr.zpusob_vzletu)
    return list(q.distinct().order_by("cas_vzletu", "id"))


def _plati(let: Let) -> str:
    return "Aeroklub" if let.plati_aeroklub else (let.platce.get_full_name() if let.platce else "–")


def souhrn(seznam: list[Let]) -> dict:
    """Součty z ukončených letů. Zrušené a probíhající lety se nepočítají."""
    ukoncene = [let for let in seznam if let.stav == StavLetu.UKONCEN]
    ucel = dict(Ucel.choices)

    def nove():
        return {"lety": 0, "minuty": 0, "tg": 0, "navijak": 0, "vlek": 0}

    def pricti(radek, let):
        radek["lety"] += 1
        radek["minuty"] += let.doba_uctovana_min or 0
        radek["tg"] += let.pocet_tg
        radek["navijak"] += let.zpusob_vzletu == ZpusobVzletu.NAVIJAK
        radek["vlek"] += let.ucel == Ucel.VLEK

    celkem = nove()
    letadla: dict[tuple, dict] = defaultdict(nove)
    platci: dict[str, dict] = defaultdict(nove)
    for let in ukoncene:
        pricti(celkem, let)
        # Vlek a ostatní lety téhož letadla vždy zvlášť (kap. 4.7 návrhu).
        pricti(letadla[(let.letadlo.imatrikulace, let.letadlo.typ, ucel[let.ucel])], let)
        pricti(platci[_plati(let)], let)
    return {
        "celkem": celkem,
        "podle_letadel": [
            {"imatrikulace": k[0], "typ": k[1], "ucel": k[2], **v}
            for k, v in sorted(letadla.items())
        ],
        "podle_platcu": [{"platce": k, **v} for k, v in sorted(platci.items())],
        "zruseno": sum(1 for let in seznam if let.stav == StavLetu.ZRUSEN),
        "neukonceno": sum(
            1 for let in seznam if let.stav in (StavLetu.VE_VZDUCHU, StavLetu.PRIPRAVEN)
        ),
    }


# --- export --------------------------------------------------------------------------


def letecky(minuty: int | None) -> str:
    return "" if minuty is None else f'{minuty // 60}°{minuty % 60}"'


SLOUPCE = [
    "Datum",
    "Letadlo",
    "Typ",
    "Kategorie",
    "Účel",
    "Úloha",
    "Způsob vzletu",
    "PIC",
    "Další posádka",
    "Hosté",
    "Platí",
    "Místo vzletu",
    "Vzlet UTC",
    "Místo přistání",
    "Přistání UTC",
    "Doba [min]",
    "Doba",
    "T&G",
    "Krátký let",
    "Vlek",
    "Soukromé",
    "Stav",
    "Zapsáno dodatečně",
    "ID letu",
]


def radky(seznam: list[Let]) -> list[list]:
    kategorie = dict(Kategorie.choices)
    ucel = dict(Ucel.choices)
    zpusob = dict(ZpusobVzletu.choices)
    funkce = dict(FunkcePosadky.choices)
    stav = dict(StavLetu.choices)
    duvod = dict(DuvodZruseni.choices)
    kratky = dict(KratkyLet.choices)
    vysledek = []
    for let in seznam:
        posadka = list(let.posadka.all())
        pic = next((p.osoba.get_full_name() for p in posadka if p.funkce == "pic"), "")
        dalsi = ", ".join(
            f"{p.osoba.get_full_name()} ({funkce[p.funkce]})" for p in posadka if p.funkce != "pic"
        )
        if let.vlecny_let_id:
            vlek = f"vlek {let.vlecny_let.letadlo.imatrikulace}"
        elif getattr(let, "vleceny_let", None):
            vlek = f"vleče {let.vleceny_let.letadlo.imatrikulace}"
        else:
            vlek = ""
        stav_text = stav[let.stav]
        if let.stav == StavLetu.ZRUSEN:
            stav_text += f" ({duvod.get(let.duvod_zruseni, '')})"
        cas = lambda d: d.astimezone(UTC).strftime("%H:%M") if d else ""  # noqa: E731
        vysledek.append(
            [
                let.cas_vzletu.astimezone(UTC).date() if let.cas_vzletu else None,
                let.letadlo.imatrikulace,
                let.letadlo.typ,
                kategorie[let.letadlo.kategorie],
                ucel[let.ucel],
                str(let.uloha) if let.uloha else "",
                zpusob[let.zpusob_vzletu],
                pic,
                dalsi,
                let.pocet_hostu or "",
                _plati(let),
                str(let.misto_vzletu),
                cas(let.cas_vzletu),
                str(let.misto_pristani) if let.misto_pristani else "",
                cas(let.cas_pristani),
                let.doba_uctovana_min,
                letecky(let.doba_uctovana_min),
                let.pocet_tg or "",
                kratky.get(let.kratky_let, ""),
                vlek,
                "ano" if let.soukrome else "",
                stav_text,
                "ano" if let.cas_pristani and let.zalozeno > let.cas_pristani else "",
                let.pk,
            ]
        )
    return vysledek


def popis_filtru(filtr: Filtr, nazvy: dict) -> list[tuple[str, str]]:
    return [
        ("Období (UTC)", f"{filtr.od:%d.%m.%Y} – {filtr.do:%d.%m.%Y}"),
        ("Letadlo", nazvy.get("letadlo") or "všechna"),
        ("Osoba v posádce", nazvy.get("osoba") or "kdokoli"),
        ("Plátce", nazvy.get("platce") or "kdokoli"),
        ("Kategorie", dict(Kategorie.choices).get(filtr.kategorie, "všechny")),
        ("Účel", dict(Ucel.choices).get(filtr.ucel, "všechny")),
        ("Způsob vzletu", dict(ZpusobVzletu.choices).get(filtr.zpusob_vzletu, "všechny")),
        ("Soukromá letadla", "včetně" if filtr.vcetne_soukromych else "bez"),
        ("Zrušené lety", "včetně" if filtr.vcetne_zrusenych else "bez"),
    ]


def _list(wb: Workbook, nazev: str, zahlavi: list[str], data: list[list]):
    ws = wb.create_sheet(nazev)
    ws.append(zahlavi)
    for bunka in ws[1]:
        bunka.font = Font(bold=True)
        bunka.fill = PatternFill("solid", fgColor="DDE7F3")
    for radek in data:
        ws.append(radek)
    for i, nadpis in enumerate(zahlavi, start=1):
        sirka = max([len(str(nadpis))] + [len(str(r[i - 1] or "")) for r in data[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(45, sirka + 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def excel(seznam: list[Let], filtr: Filtr, nazvy: dict, kdo: str) -> bytes:
    data = souhrn(seznam)
    wb = Workbook()
    wb.remove(wb.active)
    lety_ws = _list(wb, "Lety", SLOUPCE, radky(seznam))
    for bunka in lety_ws["A"][1:]:
        bunka.number_format = "DD.MM.YYYY"
    hlavicka = ["Lety", "Minuty", "Doba", "T&G", "Starty navijákem", "Vleky"]

    def hodnoty(r):
        return [r["lety"], r["minuty"], letecky(r["minuty"]), r["tg"], r["navijak"], r["vlek"]]

    _list(
        wb,
        "Podle letadel",
        ["Letadlo", "Typ", "Účel", *hlavicka],
        [[r["imatrikulace"], r["typ"], r["ucel"], *hodnoty(r)] for r in data["podle_letadel"]]
        + [["Celkem", "", "", *hodnoty(data["celkem"])]],
    )
    _list(
        wb,
        "Podle plátců",
        ["Platí", *hlavicka],
        [[r["platce"], *hodnoty(r)] for r in data["podle_platcu"]]
        + [["Celkem", *hodnoty(data["celkem"])]],
    )
    parametry = popis_filtru(filtr, nazvy) + [
        ("Vytvořeno (UTC)", timezone.now().astimezone(UTC).strftime("%d.%m.%Y %H:%M")),
        ("Vytvořil", kdo),
        ("Počítají se", "jen ukončené lety; zrušené a probíhající ne"),
        ("Neukončené lety v období", str(data["neukonceno"])),
    ]
    _list(wb, "Parametry", ["Parametr", "Hodnota"], [list(p) for p in parametry])
    vystup = io.BytesIO()
    wb.save(vystup)
    return vystup.getvalue()


def csv_data(seznam: list[Let]) -> bytes:
    """CSV pro Excel v češtině: středník a BOM (aby se správně načetla diakritika)."""
    text = io.StringIO()
    zapis = csv.writer(text, delimiter=";")
    zapis.writerow(SLOUPCE)
    for r in radky(seznam):
        zapis.writerow([(d.strftime("%d.%m.%Y") if isinstance(d, date) else d) for d in r])
    return ("﻿" + text.getvalue()).encode("utf-8")


def smi_exportovat(osoba) -> bool:
    return osoba.is_staff or osoba.role_ucetni
