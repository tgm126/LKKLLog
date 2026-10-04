from openpyxl import load_workbook
from playwright.sync_api import expect

from lety.models import StavLetu
from osoby.models import Osoba

from .conftest import HESLO, let_pilota, prihlasit


def test_ucetni_stahne_excel(page, svet, tmp_path):
    Osoba.objects.create_user(
        "ucetni@example.com", HESLO, jmeno="Filip", prijmeni="Účetní", role_ucetni=True
    )
    let_pilota(svet, stav=StavLetu.UKONCEN, minut=90)
    prihlasit(page, svet, "ucetni@example.com")
    page.get_by_role("link", name="Výpis").click()
    expect(page.get_by_role("heading", name="Výpis letů")).to_be_visible()
    expect(page.get_by_text("Lety (1)")).to_be_visible()
    expect(page.get_by_role("cell", name='0°10"').first).to_be_visible()

    with page.expect_download() as stazeni:
        page.get_by_role("link", name="Stáhnout Excel").click()
    cesta = tmp_path / stazeni.value.suggested_filename
    stazeni.value.save_as(cesta)
    assert cesta.name.startswith("lkkllog-vypis-")
    wb = load_workbook(cesta)
    assert wb.sheetnames == ["Lety", "Podle letadel", "Podle plátců", "Parametry"]
    assert wb["Lety"]["B2"].value == "OK-TCS"


def test_pilot_vidi_vypis_bez_exportu(page, svet):
    let_pilota(svet, stav=StavLetu.UKONCEN, minut=90)
    prihlasit(page, svet, "pilot@example.com")
    page.get_by_role("link", name="Výpis").click()
    expect(page.get_by_text("Lety (1)")).to_be_visible()
    expect(page.get_by_role("link", name="Stáhnout Excel")).to_have_count(0)
    # Řádek letu otevře detail s historií.
    page.get_by_role("cell", name="OK-TCS").last.click()  # řádek letu (první je v souhrnu)
    expect(page.get_by_role("dialog").get_by_text("Historie změn")).to_be_visible()
